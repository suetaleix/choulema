"""V6 routes; legacy routes remain available for migration."""
import csv
import io
import json
import zipfile
from xml.etree.ElementTree import ParseError
from typing import Optional, Literal

from fastapi import APIRouter, Request, UploadFile, File
from fastapi.responses import JSONResponse, Response, PlainTextResponse
from pydantic import BaseModel, Field

import db
import v6_commerce as commerce
import v6_payments as payments
import v6_engine as engine
import v7_referrals


class SaveIn(BaseModel):
    config: engine.DrawConfig
    id: Optional[str] = None
    revision: Optional[int] = None


class RunIn(BaseModel):
    config: engine.DrawConfig
    rehearsal: bool = False


class DrawIn(BaseModel):
    request_id: str = Field(min_length=8, max_length=100)
    control_token: str = Field(min_length=16, max_length=100)
    award_id: str = Field(default="", max_length=80)
    controller_id: str = Field(default="", max_length=100)


class TakeControlIn(BaseModel):
    controller_id: str = Field(min_length=8, max_length=100)


class ControlIn(BaseModel):
    control_token: str = Field(min_length=16, max_length=100)


class AppearanceIn(ControlIn):
    machine: str = Field(max_length=50)
    theme: str = Field(max_length=50)
    sound: str = Field(min_length=1, max_length=50)


class ForfeitIn(ControlIn):
    draw_id: int
    candidate_id: str = Field(max_length=80)
    reason: str = Field(min_length=1, max_length=200)


class CheckoutIn(BaseModel):
    slug: str = Field(max_length=50)
    request_id: str = Field(min_length=8, max_length=100)
    channel: str = Field(default="pc", max_length=10)
    price_version: int = Field(ge=1)


class ProductUpdateIn(BaseModel):
    price_cents: int
    sale_visible: bool
    reason: str


class CardBatchIn(BaseModel):
    feature: str
    days: int
    quantity: int
    label: str
    purpose: str = ""
    redeem_deadline: Optional[int] = None


class CardRedeemIn(BaseModel):
    code: str


class CardVoidIn(BaseModel):
    reason: str


def install(app, current_subject):
    router = APIRouter(prefix="/api/v6")

    def owner(req):
        subject = current_subject(req)
        if not subject:
            raise engine.DomainError("请刷新重试", 401)
        key = f"{subject['kind']}:{subject['id']}"
        expected = req.headers.get("X-Draft-Scope")
        if expected and expected != key:
            raise engine.DomainError("登录账号已变化，请刷新后继续", 409)
        return key

    def admin(req):
        subject = current_subject(req)
        if not subject or subject["kind"] != "user" or subject["user"]["role"] != "superadmin":
            raise engine.DomainError("无权限", 403)
        return subject["id"]

    @app.exception_handler(engine.DomainError)
    async def domain_error(req, exc):
        return JSONResponse({"error": exc.message}, status_code=exc.status)

    @router.get("/status")
    def status(req: Request):
        subject = current_subject(req)
        key = owner(req)
        with db.conn() as c:
            handoffs = [r[0] for r in c.execute("SELECT guest_scope FROM v6_guest_handoffs WHERE owner=? ORDER BY created_at, rowid", (key,))] if subject["kind"] == "user" else []
        return {"kind": subject["kind"], "draft_scope": key, "guest_draft_scopes": handoffs,
                "name": (subject.get("user") or {}).get("email") or (subject.get("user") or {}).get("custom_id", "游客"),
                "role": (subject.get("user") or {}).get("role"),
                "features": sorted(engine.features(key)), "capabilities": [c for c in engine.content_catalog.CONTENT["capabilities"] if engine.has_capability(key, c)], "content": engine.content_catalog.public_catalog(), "version": 6}

    @router.get("/configs")
    def configs(req: Request):
        key = owner(req)
        legacy = []
        if key.startswith("user:"):
            for tpl in db.list_templates(int(key.split(":")[1])):
                if not tpl.get("items"):
                    continue
                app_kind = {"dish": "food", "person": "person", "prize": "prize"}.get(tpl["type"], "food")
                recommended = any(
                    tpl["title"] == preset["title"] and tpl["type"] == preset["type"]
                    and [(i["name"], i.get("weight", 1)) for i in tpl["items"]]
                    == [(i["name"], i.get("weight", 1)) for i in preset["items"]]
                    for preset in db.DEFAULT_TEMPLATES
                )
                legacy.append({"legacy_id": tpl["id"], "recommended": recommended, "config": {"title": tpl["title"], "app": app_kind,
                    "mode": "no_repeat" if app_kind == "person" else "normal", "machine": "sphere" if app_kind == "person" else "capsule", "theme": "classic", "count": 1,
                    "items": [{"id": str(i["id"]), "name": i["name"], "weight": max(1, i.get("weight", 1)), "stock": 1} for i in tpl["items"]]}})
        return {"configs": engine.list_configs(key), "legacy": legacy}

    @router.post("/import")
    async def import_roster(req: Request, file: UploadFile = File(...)):
        if not engine.has_capability(owner(req), "roster_import"):
            raise engine.DomainError("文件导入需要专业权益；基础版可直接粘贴名单", 403)
        data = await file.read(2_000_001)
        if len(data) > 2_000_000:
            raise engine.DomainError("文件请控制在 2 MB 内")
        name = (file.filename or "").lower()
        try:
            if name.endswith(".csv"):
                rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
            elif name.endswith(".xlsx"):
                from openpyxl import load_workbook
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    if sum(i.file_size for i in z.infolist()) > 20_000_000:
                        raise ValueError("expanded workbook too large")
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                try:
                    rows = []
                    for row in workbook.active.iter_rows(values_only=True):
                        rows.append(list(row[:20]))
                        if len(rows) > 1001:
                            raise ValueError("too many rows")
                finally:
                    workbook.close()
            else:
                raise engine.DomainError("支持 UTF-8 CSV 或 XLSX 文件")
            if not rows or len(rows) > 1001:
                raise ValueError("empty or too many rows")
            # Return a preview so users explicitly map fields before replacing a roster.
            clean = [[str(value if value is not None else "")[:100] for value in row[:20]] for row in rows]
        except (ValueError, UnicodeError, OSError, zipfile.BadZipFile, KeyError, ParseError):
            raise engine.DomainError("文件无法读取，请检查格式；最多 1000 名候选")
        return {"rows": clean}

    @router.get("/configs/{cid}/report")
    def class_report(cid: str, req: Request):
        key = owner(req)
        if not engine.has_capability(key, "class_report"):
            raise engine.DomainError("跨课次统计需要专业权益", 403)
        saved = next((r for r in engine.list_configs(key) if r["id"] == cid), None)
        if not saved:
            raise engine.DomainError("配置不存在", 404)
        records = {}
        sessions = 0
        with db.conn() as c:
            runs = c.execute("SELECT id,payload FROM v6_runs WHERE owner=? AND rehearsal=0", (key,)).fetchall()
            for r in runs:
                cfg = json.loads(r["payload"])
                if cfg.get("source_config_id") != cid:
                    continue
                sessions += 1
                for d in c.execute("SELECT result FROM v6_draws WHERE run_id=?", (r["id"],)):
                    for item in json.loads(d["result"])["results"]:
                        entry = records.setdefault(item["id"], {"id": item["id"], "name": item["name"], "count": 0})
                        entry["count"] += 1
        return {"title": saved["config"]["title"], "sessions": sessions, "students": list(records.values())}

    @router.get("/classes")
    def classes(req: Request):
        key = owner(req)
        if not engine.has_capability(key, "class_sessions"):
            raise engine.DomainError("班级快捷切换需要专业权益；基础版可在我的配置载入班级", 403)
        saved = {r["id"]: {"id": r["id"], "title": r["config"]["title"], "revision": r["revision"], "sessions": 0, "active_run_id": None, "pending_changes": False}
                 for r in engine.list_configs(key) if r["config"]["app"] == "person"}
        with db.conn() as c:
            for run in c.execute("SELECT id,payload,archived FROM v6_runs WHERE owner=? AND rehearsal=0 ORDER BY created_at DESC,rowid DESC", (key,)):
                cfg = json.loads(run["payload"])
                target = saved.get(cfg.get("source_config_id"))
                if not target or cfg["app"] != "person":
                    continue
                target["sessions"] += 1
                if not run["archived"] and not target["active_run_id"]:
                    target["active_run_id"] = run["id"]
                    target["pending_changes"] = cfg.get("source_config_revision", 0) != target["revision"]
        return {"classes": list(saved.values())}

    @router.post("/configs")
    def save(body: SaveIn, req: Request):
        return engine.save_config(owner(req), body.config, body.id, body.revision)

    @router.post("/runs")
    def create(body: RunIn, req: Request):
        return engine.new_run(owner(req), body.config, body.rehearsal)

    @router.get("/runs")
    def runs(req: Request):
        return {"runs": engine.list_runs(owner(req))}

    @router.get("/runs/{rid}")
    def run(rid: str, req: Request):
        return engine.get_run(owner(req), rid)

    @router.post("/runs/{rid}/draw")
    def draw(rid: str, body: DrawIn, req: Request):
        key = owner(req)
        result = engine.draw(key, rid, body.request_id, body.control_token, body.award_id, body.controller_id)
        if key.startswith("user:"):
            v7_referrals.qualify_user(int(key.split(":",1)[1]))
        return result

    @router.post("/runs/{rid}/start")
    def start(rid: str, body: DrawIn, req: Request):
        return engine.start(owner(req), rid, body.request_id, body.control_token, body.award_id, body.controller_id)

    @router.post("/runs/{rid}/control")
    def take_control(rid: str, body: TakeControlIn, req: Request):
        return engine.take_control(owner(req), rid, body.controller_id)

    @router.post("/runs/{rid}/cancel")
    def cancel(rid: str, body: DrawIn, req: Request):
        engine.cancel_presentation(owner(req), rid, body.request_id, body.control_token, body.controller_id)
        return {"ok": True}

    @router.post("/runs/{rid}/archive")
    def archive(rid: str, body: ControlIn, req: Request):
        engine.archive(owner(req), rid, body.control_token)
        return {"ok": True}

    @router.post("/runs/{rid}/appearance")
    def appearance(rid: str, body: AppearanceIn, req: Request):
        return engine.change_appearance(owner(req), rid, body.control_token, body.machine, body.theme, body.sound)

    @router.post("/runs/{rid}/forfeit")
    def forfeit(rid: str, body: ForfeitIn, req: Request):
        engine.forfeit(owner(req), rid, body.control_token, body.draw_id, body.candidate_id, body.reason)
        return {"ok": True}

    @router.get("/runs/{rid}/export")
    def export(rid: str, req: Request, award_id: str = ""):
        key = owner(req)
        if not engine.has_capability(key, "result_export"):
            raise engine.DomainError("按场次导出需要专业权益；当前结果可免费复制", 403)
        state = engine.get_run(key, rid)
        if award_id and award_id not in {a["id"] for a in state["config"]["awards"]}:
            raise engine.DomainError("奖项不存在", 404)
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(["场次", "演练", "奖项", "编号", "姓名/结果", "时间", "状态", "抽取记录编号", "弃奖原因", "弃奖时间", "小组"])
        forfeited = {(event["payload"]["draw_id"], event["payload"]["candidate_id"]): event for event in state["events"] if event["kind"] == "forfeit"}
        def safe(v):
            v = str(v)
            return "'" + v if v.lstrip().startswith(("=", "+", "-", "@")) or v.startswith(("\t", "\r", "\n")) else v
        for d in state["draws"]:
            if award_id and d["award_id"] != award_id:
                continue
            groups = {item["id"]: index + 1 for index, group in enumerate(d.get("groups", [])) for item in group}
            for item in d["results"]:
                event = forfeited.get((d["id"], item["id"]))
                writer.writerow([safe(state["config"]["title"]), state["rehearsal"], safe(d["award"]), safe(item["id"]), safe(item["name"]), d["created_at"], "弃奖" if event else "有效", d["id"], safe(event["payload"]["reason"]) if event else "", event["created_at"] if event else "", groups.get(item["id"], "")])
        return Response("\ufeff" + buf.getvalue(), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="choulema-results.csv"'})

    @router.get("/display/{rid}")
    def display(rid: str, request: Request):
        # Token passed in a header, not logged in query strings.
        return engine.view_run(rid, request.headers.get("X-View-Token", ""))

    @router.get("/shop")
    def shop(req: Request):
        return {"products": payments.catalog(owner(req))}

    @router.get("/account")
    def account(req: Request):
        key = owner(req)
        if not key.startswith("user:"):
            raise engine.DomainError("请先免费注册或登录", 401)
        return commerce.entitlement_details(key)

    @router.get("/admin/products")
    def admin_products(req: Request):
        admin(req)
        return {"products": payments.catalog(owner(req))}

    @router.post("/admin/products/{slug}")
    def admin_product_update(slug: str, body: ProductUpdateIn, req: Request):
        actor = admin(req)
        try:
            return commerce.set_product(slug, body.price_cents, body.sale_visible, actor, body.reason, payments.CATALOG)
        except ValueError as exc:
            raise engine.DomainError(str(exc), 400) from exc

    @router.get("/admin/products/{slug}/history")
    def admin_product_history(slug: str, req: Request):
        admin(req)
        try:
            return {"versions": commerce.price_history(slug, payments.CATALOG)}
        except ValueError as exc:
            raise engine.DomainError(str(exc), 404) from exc

    @router.post("/cards/redeem")
    def redeem_card(body: CardRedeemIn, req: Request):
        key = owner(req)
        if not key.startswith("user:"):
            raise engine.DomainError("请先免费注册或登录", 401)
        try:
            return commerce.redeem_card(body.code, int(key.split(":", 1)[1]))
        except ValueError as exc:
            raise engine.DomainError(str(exc), 400) from exc

    @router.get("/admin/card-batches")
    def card_batches(req: Request):
        admin(req)
        return {"batches": commerce.list_card_batches()}

    @router.get("/admin/card-batches/{batch_id}/cards")
    def batch_cards(batch_id: str, req: Request):
        admin(req)
        return {"cards": commerce.list_batch_cards(batch_id)}

    @router.post("/admin/card-batches")
    def create_card_batch(body: CardBatchIn, req: Request):
        actor = admin(req)
        try:
            return commerce.generate_cards(body.feature, body.days, body.quantity, body.label,
                                           body.purpose, body.redeem_deadline, actor, payments.CATALOG)
        except ValueError as exc:
            raise engine.DomainError(str(exc), 400) from exc

    @router.post("/admin/cards/{card_id}/void")
    def void_card(card_id: str, body: CardVoidIn, req: Request):
        actor = admin(req)
        try:
            if not commerce.void_card(card_id, actor, body.reason):
                raise engine.DomainError("卡密不存在或已使用／作废", 409)
        except ValueError as exc:
            raise engine.DomainError(str(exc), 400) from exc
        return {"success": True}

    @router.post("/checkout")
    def checkout(body: CheckoutIn, req: Request):
        return payments.checkout(owner(req), body.slug, body.request_id, body.channel, body.price_version)

    @router.get("/admin/orders")
    def admin_orders(req: Request, q: str = "", slug: str = "", state: str = "",
                     provider: str = "", channel: str = "", issue: str = "", from_ts: Optional[int] = None,
                     to_ts: Optional[int] = None, page: int = 1, page_size: int = 20):
        admin(req)
        return payments.admin_orders(q=q, slug=slug, state=state, provider=provider,
                                     channel=channel, issue=issue, from_ts=from_ts, to_ts=to_ts,
                                     page=page, page_size=page_size)

    @router.post("/admin/orders/{number}/reconcile")
    def admin_reconcile_order(number: str, req: Request):
        admin(req)
        return payments.reconcile(number)

    @router.get("/admin/overview")
    def admin_overview(req: Request):
        admin(req)
        return {**commerce.admin_overview(), "order_issues": payments.admin_issue_counts()}

    @router.get("/admin/audit")
    def admin_audit(req: Request):
        admin(req)
        return {"events": commerce.admin_audit()}

    @router.get("/admin/users/{uid}/commerce")
    def admin_user_commerce(uid: int, req: Request):
        admin(req)
        user = db.get_user(uid)
        if not user:
            raise engine.DomainError("账号不存在", 404)
        owner_key = f"user:{uid}"
        return {"user": db.user_public(user), "entitlements": commerce.entitlement_details(owner_key),
                "orders": payments.list_orders(owner_key)}

    @router.get("/orders")
    def orders(req: Request):
        return {"orders": payments.list_orders(owner(req))}

    @router.get("/orders/{number}")
    def order(number: str, req: Request):
        return payments.order(owner(req), number)

    @router.post("/orders/{number}/refresh")
    def refresh_order(number: str, req: Request):
        return payments.refresh(owner(req), number)

    @router.post("/payment/notify")
    async def notify(req: Request):
        try:
            if "application/x-www-form-urlencoded" not in req.headers.get("content-type", ""):
                raise engine.DomainError("通知格式无效", 400)
            pairs = (await req.form()).multi_items()
            data = {}
            for key, value in pairs:
                if key in data:
                    raise engine.DomainError("通知参数重复", 400)
                data[key] = value
            payments.verify_notification(data)
            payments.settle(data)
        except (engine.DomainError, ValueError, TypeError):
            return PlainTextResponse("fail", status_code=400)
        return PlainTextResponse("success")

    app.include_router(router)
